# اختبار قسم الإحصائيات = كشف الخدمات اليومية (واجهة + طباعة PDF حقيقي + تنزيل DOCX حقيقي وتحليله).
# يعمل بمحاكاة Firebase في الذاكرة (dev/mockfb.py) لأن بيئة الاختبار لا تصل إلى Firebase الحقيقي؛
# الكود نفسه لا يحتوي أي محاكاة. السيرفر على المنفذ 8099.
import os, re, sys, subprocess, unicodedata, zipfile
sys.path.insert(0, os.path.dirname(__file__))
from datetime import date, timedelta
from playwright.sync_api import sync_playwright
from pypdf import PdfReader
import docx as pydocx
from lxml import etree
from mockfb import APP, AUTH, FS

BASE = "http://localhost:8099/"; V = "https://www.gstatic.com/firebasejs/10.14.1/"; OUT = "/tmp/smoke"
errors, results = [], []
def check(name, cond, extra=""):
    results.append(bool(cond)); print(("PASS " if cond else "FAIL ") + name + (f"  [{extra}]" if extra else ""))
def setup(b, w=1280, h=900):
    ctx = b.new_context(viewport={"width": w, "height": h}, accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(f"pageerror: {e}"))
    pg.on("console", lambda m: errors.append(f"console: {m.text}") if m.type == "error" else None)
    for n, body in (("firebase-app.js", APP), ("firebase-auth.js", AUTH), ("firebase-firestore.js", FS)):
        pg.route(V + n, lambda r, _q, body=body: r.fulfill(status=200, content_type="application/javascript", body=body))
    pg.route(re.compile(r"https://fonts\..*"), lambda r, _q: r.fulfill(status=200, content_type="text/css", body=""))
    return pg
def store(pg): return pg.evaluate("JSON.parse(localStorage.getItem('mockstore')||'{}')")
def raw(pg): return pg.evaluate("localStorage.getItem('mockstore')")
def inject(pg, patch):
    pg.evaluate("p=>{const s=JSON.parse(localStorage.getItem('mockstore'));for(const [k,v] of Object.entries(p)){if(v===null)delete s[k];else s[k]=v}localStorage.setItem('mockstore',JSON.stringify(s))}", patch)
def login(pg):
    pg.goto(BASE); pg.wait_for_selector("#login-screen:not([hidden])")
    pg.fill("#login-email", "me@test.com"); pg.fill("#login-password", "secret12"); pg.click("#login-btn")
    pg.wait_for_selector("#app-shell:not([hidden])"); pg.wait_for_selector(".toast-ok")
def go(pg, route, sel): pg.click(f'a[data-route="{route}"]'); pg.wait_for_selector(sel)
def pid(n): return f"p_seed_{n:02d}"
def fmt(i): y, m, d = i.split("-"); return f"{d}/{m}/{y}"
def iso(n): return (T + timedelta(days=n)).isoformat()
def D(d, **o): return {"date": d, "status": "draft", "parmanoPersonId": None, "gardPersonId": None, "restPersonIds": [], "otherTasks": [], **o}
def period(pg, a, b_):
    pg.fill("#rp-from", a); pg.fill("#rp-to", b_); pg.dispatch_event("#rp-to", "change")
def apply(pg):
    pg.click("#rp-apply"); pg.wait_for_function("!document.getElementById('rp-apply').disabled")
def rows(pg): return pg.locator("#rp-table tbody tr.rp-row")
def row(pg, n): return pg.locator(f"#rp-table tbody tr[data-date='{iso(n)}']")
def lines(pg, n, col): return [l.strip() for l in row(pg, n).locator(f"td.rp-c-{col} .rp-line").all_inner_texts()]
def kind(pg, n): return re.search(r"rp-(done|pending|missing|upcoming|today)", row(pg, n).get_attribute("class")).group(1)
def ptext(path, n): return unicodedata.normalize("NFKC", subprocess.run(["pdftotext", "-f", str(n), "-l", str(n), "-layout", path, "-"], capture_output=True, text=True).stdout)

T = date.today()
with sync_playwright() as p:
    b = p.chromium.launch(); pg = setup(b); login(pg)
    go(pg, "people", ".person")
    inject(pg, {
        f"daily_services/{iso(-4)}": D(iso(-4), status="final", parmanoPersonId=pid(1), gardPersonId=pid(2), restPersonIds=[pid(3), pid(7), pid(8)], otherTasks=[{"personId": pid(11), "task": "حراسة المخزن"}]),
        f"daily_services/{iso(-3)}": D(iso(-3), status="final", parmanoPersonId=pid(4), gardPersonId=pid(5)),
        f"daily_services/{iso(-2)}": D(iso(-2), parmanoPersonId=pid(9)),
        f"daily_services/{iso(0)}":  D(iso(0), parmanoPersonId=pid(1)),
        f"daily_services/{iso(1)}":  D(iso(1), status="final", parmanoPersonId=pid(2), gardPersonId=pid(1)),
        f"daily_services/{iso(2)}":  D(iso(2), restPersonIds=[pid(3), pid(7), pid(8), pid(10)]),
        f"daily_services/{iso(4)}":  D(iso(4), parmanoPersonId=pid(12), restPersonIds=[pid(16)]),
        "assignments/a_leave": {"personId": pid(6), "type": "إجازة", "startDate": iso(-1), "endDate": iso(1), "notes": ""},
        "assignments/a_open": {"personId": pid(13), "type": "مهمة أخرى", "startDate": iso(-2), "endDate": None, "notes": ""},
        f"people/{pid(16)}": {**store(pg)[f"people/{pid(16)}"], "isActive": False},
    })
    base_raw = raw(pg)
    go(pg, "stats", "#rp-table")

    # ===== 1) القسم = الكشف فقط (لا إحصائيات قديمة) =====
    body = pg.inner_text("#view")
    check("عنوان الصفحة «الإحصائيات» والتنقل", pg.inner_text("#page-title") == "الإحصائيات" and pg.get_attribute('a[data-route="stats"]', "aria-current") == "page")
    check("المحتوى كشف الخدمات اليومية فقط", "كشف الخدمات اليومية" in body and pg.locator("#rp-table").count() == 1)
    check("حُذفت الإحصائيات القديمة (الملخص الشهري/السنوي/المقارنة/البطاقات/السجل)", not any(x in body for x in ("الملخص الشهري", "الملخص السنوي", "مقارنة الأفراد", "مقارنة بالمتوسط", "متوسط الفاصل", "يوم معتمد", "مسودة (غير محتسبة)")) and pg.locator(".stat-card, .stat-cards, .bar, #st-compare, #st-monthly, #st-yearly").count() == 0)
    # الفترة الافتراضية: الشهر الحالي
    first = T.replace(day=1); last = (T.replace(day=28) + timedelta(days=4)); last = last - timedelta(days=last.day)
    check("الفترة الافتراضية = الشهر الحالي وعدد أيامه", pg.input_value("#rp-from") == first.isoformat() and pg.input_value("#rp-to") == last.isoformat() and str((last - first).days + 1) in pg.inner_text("#rp-days"), pg.inner_text("#rp-days"))
    check("واجهة الفترة: بداية، نهاية، «عرض الكشف»، البحث، نوع الخدمة، إعادة العرض", all(pg.locator(s).count() == 1 for s in ("#rp-from", "#rp-to", "#rp-apply", "#rp-search", "#rp-type", "#rp-reset")) and pg.inner_text("#rp-apply").strip() == "عرض الكشف")
    check("خيارات النوع: جميع الخدمات (الافتراضي) + الأربعة", pg.locator("#rp-type option").all_inner_texts() == ["جميع الخدمات", "الخفارة", "الحراسة", "الراحة", "المهام"] and pg.input_value("#rp-type") == "all")

    # ===== 2) فترة تشمل ماضيًا وقادمًا =====
    period(pg, iso(-4), iso(4)); check("عدد الأيام المشمولة يتحدّث فورًا (9 أيام)", "9 أيام" in pg.inner_text("#rp-days"))
    check("تنبيه أن الفترة لم تُطبَّق بعد", "اضغط «عرض الكشف»" in pg.inner_text("#rp-msg"))
    apply(pg)
    check("بعد التطبيق: صف لكل يوم مرتبة زمنيًا (9)", [r.get_attribute("data-date") for r in rows(pg).all()] == [iso(i) for i in range(-4, 5)])
    check("لا تنبيه بعد التطبيق، وترويسة الفترة", pg.inner_text("#rp-msg").strip() == "" and f"من {fmt(iso(-4))} إلى {fmt(iso(4))}" in pg.inner_text(".rp-period").replace("\n", " "))
    check("أعمدة ثابتة بالترتيب المطلوب", pg.locator("#rp-table thead th").all_inner_texts() == ["التاريخ واليوم", "الخفارة", "الحراسة", "الراحة", "المهام"] and all(r.locator("td").count() == 5 for r in rows(pg).all()))
    check("حالات الأيام صحيحة (منفذ/منفذ/غير محسوم/غير مُعدّ/اليوم/قادم×4)", [kind(pg, i) for i in range(-4, 5)] == ["done", "done", "pending", "missing", "today", "upcoming", "upcoming", "upcoming", "upcoming"], str([kind(pg, i) for i in range(-4, 5)]))
    st = lambda n: row(pg, n).locator(".rp-state").inner_text().strip()
    check("نص «منفذ» لليومين المعتمدين الماضيين فقط", st(-4) == "✓ منفذ" and st(-3) == "✓ منفذ" and "منفذ" not in " ".join(st(i) for i in range(-2, 5)))
    check("يوم قادم معتمد مسبقًا يُوصف «قادم (معتمد)» لا منفذ", st(1) == "قادم (معتمد)" and st(2) == "قادم (مسودة)" and st(3) == "قادم (غير مُعدّ)")
    check("اليوم الحالي موسوم بحالته، والماضي غير المحسوم والمفقود بتسمية صريحة", st(0) == "اليوم (مسودة)" and st(-2) == "مسودة غير محسومة" and st(-1) == "غير مُعدّ")
    check("اسم يوم الأسبوع بجانب التاريخ", row(pg, -4).locator(".rp-wd").inner_text().strip() in ("الأحد", "الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت") and row(pg, -4).locator(".rp-d").inner_text().strip() == fmt(iso(-4)))
    check("وسيلة الإيضاح: 5 حالات بنصوصها", pg.locator(".rp-legend li").count() == 5 and all(x in pg.inner_text(".rp-legend") for x in ("منفذ", "قادم", "اليوم الحالي", "مسودة غير محسومة", "غير مُعدّ")))

    pg.screenshot(path=f"{OUT}/rp_desktop.png", full_page=True)
    # ===== 3) الألوان والتمييز البصري =====
    bg = lambda n: row(pg, n).locator("td.rp-cell").first.evaluate("e=>getComputedStyle(e).backgroundColor")
    check("المنفذ: خلفية محايدة؛ القادم: شبه أبيض أخفّ؛ غير المحسوم: خافتة", bg(-4) == "rgb(236, 239, 242)" and bg(1) == "rgb(253, 253, 254)" and bg(-2) == "rgb(255, 247, 224)", f"{bg(-4)} {bg(1)} {bg(-2)}")
    check("اليوم الحالي: خط جانبي ثابت في خلية التاريخ", "inset" in row(pg, 0).locator("td.rp-date").evaluate("e=>getComputedStyle(e).boxShadow") and "inset" not in row(pg, 1).locator("td.rp-date").evaluate("e=>getComputedStyle(e).boxShadow"))
    check("لا ألوان مشبعة: كل الخلفيات قريبة من الأبيض", all(sum(map(int, re.findall(r"\d+", bg(i))[:3])) > 700 for i in range(-4, 5)))

    # ===== 4) محتوى الخلايا =====
    check("أسماء الخفارة والحراسة والراحة في أعمدتها وكل اسم في سطر مستقل", lines(pg, -4, "parmano") == ["بن ناصر"] and lines(pg, -4, "gard") == ["عقون"] and lines(pg, -4, "rest") == ["بوثلجة", "حميدي", "دبزة"], str(lines(pg, -4, "rest")))
    check("المهمة اليدوية باسم الفرد والمهمة، والسائق 🚗", lines(pg, -4, "tasks") == ["مرزوقي 🚗 — حراسة المخزن"], str(lines(pg, -4, "tasks")))
    check("لا خلط: الخفير ليس في عمود الحراسة والعكس", all("بن ناصر" not in l for l in lines(pg, -4, "gard")) and all("عقون" not in l for l in lines(pg, -4, "parmano")))
    check("البديل بسطر مستقل وبنوعه (اليوم الثاني)", lines(pg, -3, "parmano") == ["فقير 🚗", "بديل: بن ناصر"] and lines(pg, -3, "gard") == ["خدام 🚗", "بديل: عقون"], str(lines(pg, -3, "parmano")))
    check("إذا لم توجد خدمة تظهر «لا توجد»", row(pg, -1).locator("td.rp-c-parmano").inner_text().strip() == "لا توجد" and row(pg, 3).locator("td.rp-c-rest").inner_text().strip() == "لا توجد" and row(pg, -4).locator("td.rp-c-tasks .rp-none").count() == 0)
    check("أربعة أفراد راحة: كلهم ظاهرون بلا قصّ، كل اسم في سطر", lines(pg, 2, "rest") == ["بوثلجة", "حميدي", "دبزة", "شيتور"])
    check("التكليفات تظهر في المهام بنوعها ومدتها (مائلة) حتى في يوم غير مُعدّ", any("غربي — إجازة" in l for l in lines(pg, -1, "tasks")) and "من " + fmt(iso(-1)) + " إلى " + fmt(iso(1)) in row(pg, -1).locator("td.rp-c-tasks").inner_text() and row(pg, -1).locator("td.rp-c-tasks .rp-assignment").count() >= 1)
    check("المهمة المستمرة تُعرض «بلا تاريخ نهاية» دون تاريخ وهمي", "بلا تاريخ نهاية" in row(pg, 3).locator("td.rp-c-tasks").inner_text() and "9999" not in pg.inner_text("#view"))
    check("الفرد المعطّل يُعلَّم ولا يُحذف اسمه", "(غير نشط)" in row(pg, 4).locator("td.rp-c-rest").inner_text() and "شيباني" in row(pg, 4).locator("td.rp-c-rest").inner_text())
    check("لا أسماء افتراضية: كل الأسماء المعروضة من قائمة الأفراد", set(re.sub(r"( 🚗| \(غير نشط\))", "", l.replace("بديل: ", "")).split(" — ")[0] for l in pg.locator("#rp-table .rp-line").all_inner_texts() for l in [l.split("\n")[0]]) <= {v["name"] for k, v in store(pg).items() if k.startswith("people/")})
    check("سطر التعداد: 9 أيام بالحالات (لا بطاقات)", "9 أيام" in pg.inner_text(".rp-counts") and "2 منفذ" in pg.inner_text(".rp-counts") and "4 قادم" in pg.inner_text(".rp-counts") and "1 اليوم الحالي" in pg.inner_text(".rp-counts"), pg.inner_text(".rp-counts"))

    # ===== 5) التحقق من الفترة =====
    period(pg, iso(5), iso(-5)); apply(pg)
    check("بداية بعد نهاية: رسالة واضحة والكشف السابق باقٍ", "تاريخ البداية بعد تاريخ النهاية" in pg.inner_text("#rp-msg") and rows(pg).count() == 9 and pg.inner_text("#rp-days").strip() == "—")
    pg.fill("#rp-from", ""); apply(pg); check("تاريخ فارغ: رسالة", "حدّد تاريخ البداية" in pg.inner_text("#rp-msg"))
    period(pg, "2025-01-01", "2026-12-31"); apply(pg); check("فترة أطول من 366 يومًا: رسالة بالحد", "الحد الأقصى 366" in pg.inner_text("#rp-msg") and rows(pg).count() == 9)
    period(pg, iso(-4), iso(4)); apply(pg); check("تصحيح الفترة يعيد العرض", rows(pg).count() == 9 and pg.inner_text("#rp-msg").strip() == "")

    # ===== 6) البحث والتصفية =====
    pg.fill("#rp-search", "بن ناصر")
    check("البحث: الأيام التي يرتبط بها الفرد فقط (5)", [r.get_attribute("data-date") for r in rows(pg).all()] == [iso(i) for i in (-4, -3, 0, 1, 2)], str([r.get_attribute("data-date") for r in rows(pg).all()]))
    check("الاسم المطابق معلَّم وبقية الأسماء ظاهرة", pg.locator("mark.rp-hit").count() >= 5 and "عقون" in row(pg, -4).inner_text() and all(m.inner_text().strip() == "بن ناصر" for m in pg.locator("mark.rp-hit").all()))
    check("التعداد يوضح المعروض من الفترة والتصفية المطبقة", "المعروض 5 من 9 أيام" in pg.inner_text(".rp-counts") and "بحث عن: «بن ناصر»" in pg.inner_text(".rp-filter") and "يطابقان التصفية" in pg.inner_text("#rp-export-note"))
    pg.fill("#rp-search", "احمد"); check("بحث بلا نتيجة: رسالة واضحة وليس جدولًا فارغًا", pg.locator(".rp-empty").count() == 1 and rows(pg).count() == 0)
    pg.fill("#rp-search", "بوثلجه"); check("بحث متسامح مع الهمزة والتاء المربوطة (بوثلجة)", rows(pg).count() == 2)
    pg.fill("#rp-search", "")
    pg.select_option("#rp-type", "rest")
    check("تصفية «الراحة»: عمود التاريخ + الراحة فقط وكل الأيام باقية", pg.locator("#rp-table thead th").all_inner_texts() == ["التاريخ واليوم", "الراحة"] and rows(pg).count() == 9 and row(pg, -4).locator("td").count() == 2)
    pg.fill("#rp-search", "بوثلجة"); check("بحث + نوع: الراحة وحدها (أيام -4 و+2)", [r.get_attribute("data-date") for r in rows(pg).all()] == [iso(-4), iso(2)])
    pg.fill("#rp-search", "بن ناصر"); check("بن ناصر ليس في الراحة: لا نتائج مع نوع الراحة", rows(pg).count() == 0)
    pg.select_option("#rp-type", "gard"); check("نوع الحراسة + بن ناصر: يوم +1 وبديل الحراسة +2", [r.get_attribute("data-date") for r in rows(pg).all()] == [iso(1), iso(2)])
    pg.click("#rp-reset")
    check("إعادة العرض الكامل: 5 أعمدة و9 أيام وبحث فارغ ونوع «جميع الخدمات»", pg.locator("#rp-table thead th").count() == 5 and rows(pg).count() == 9 and pg.input_value("#rp-search") == "" and pg.input_value("#rp-type") == "all" and pg.locator(".rp-filter").count() == 0)

    # ===== 7) الوصول إلى التعديل الحالي دون أي كتابة =====
    check("رابط «فتح اليوم» لكل يوم يشير إلى صفحة الخدمات الحالية", all(r.locator("a.rp-link").get_attribute("href") == f"#/services/{r.get_attribute('data-date')}" for r in rows(pg).all()))
    check("الكشف للقراءة فقط: لا نقر على خلية يكتب شيئًا", (row(pg, 2).locator("td.rp-c-rest").click(), raw(pg) == base_raw)[1])
    row(pg, 2).locator("a.rp-link").click(); pg.wait_for_selector(".day-card")
    check("«فتح اليوم» يفتح اليوم المعني في محرر الخدمات الحالي", pg.inner_text("#page-title") == "برنامج الخدمات" and fmt(iso(2)) in pg.inner_text(".day-title") and pg.locator("[data-action=pick-rest]").count() == 1)
    check("لم يُكتب شيء في Firebase (المحاكاة) بمجرد العرض والتصفية والفتح", raw(pg) == base_raw)
    go(pg, "stats", "#rp-table")

    # ===== 8) الطباعة (وضع الطباعة) =====
    period(pg, iso(-4), iso(4)); apply(pg)
    pg.evaluate("()=>{window.__p={n:0,titles:[]};window.print=()=>{window.__p.n++;window.__p.titles.push(document.title);window.dispatchEvent(new Event('afterprint'))}}")
    t0 = pg.title(); pg.click("[data-action=print-report]")
    check("«طباعة الكشف» يفتح نافذة الطباعة مرة واحدة بعنوان يحمل الفترة ثم يُستعاد", pg.evaluate("window.__p.n") == 1 and f"{iso(-4)} إلى {iso(4)}" in pg.evaluate("window.__p.titles[0]") and pg.title() == t0)
    pg.click("[data-action=save-pdf]")
    check("«حفظ بصيغة PDF» يمر عبر نافذة الطباعة مع إرشاد «حفظ كـ PDF»", pg.evaluate("window.__p.n") == 2 and pg.locator(".toast:has-text('حفظ كـ PDF')").count() >= 1)
    pg.set_viewport_size({"width": 794, "height": 1123}); pg.emulate_media(media="print")
    vis = lambda s: pg.is_visible(s)
    check("في الطباعة: الأدوات والتنقل والروابط مخفية والكشف ظاهر", not any(vis(s) for s in (".rp-controls", ".side", ".topbar", ".toasts")) and vis("#rp-table") and vis(".rp-legend") and pg.locator(".rp-link").first.is_hidden())
    check("في الطباعة: جدول حقيقي بخمسة أعمدة ورأس مكرر وكل الأيام التسعة", pg.evaluate("[getComputedStyle(document.querySelector('#rp-table')).display,getComputedStyle(document.querySelector('#rp-table thead')).display]") == ["table", "table-header-group"] and rows(pg).count() == 9 and pg.locator("#rp-table thead th").count() == 5)
    check("منع انقسام الصف بين صفحتين", pg.evaluate("getComputedStyle(document.querySelector('#rp-table tbody tr')).breakInside") == "avoid")
    check("العنوان والفترة وتاريخ الإعداد والوسيلة في الطباعة", all(x in pg.inner_text("#rp-report") for x in ("كشف الخدمات اليومية", f"من {fmt(iso(-4))} إلى {fmt(iso(4))}", "تاريخ إعداد التقرير: " + fmt(T.isoformat()), "منفذ — يوم مضى")))
    colors = pg.evaluate("""()=>[...document.querySelectorAll('#rp-report .rp-line, #rp-report .rp-d, #rp-report .rp-state, #rp-report th')].filter(e=>e.getClientRects().length).map(e=>getComputedStyle(e).color)""")
    check("الطباعة بالأبيض والأسود: كل النصوص سوداء", colors and all(c == "rgb(0, 0, 0)" for c in colors), str(set(colors)))
    bgp = lambda n: row(pg, n).locator("td.rp-cell").first.evaluate("e=>getComputedStyle(e).backgroundColor")
    check("المنفذ رمادي فاتح جدًا، والقادم أبيض، واليوم بحدود سميكة سوداء", bgp(-4) == "rgb(236, 239, 241)" and bgp(1) == "rgb(255, 255, 255)" and float(row(pg, 0).locator("td").first.evaluate("e=>getComputedStyle(e).borderTopWidth").replace("px", "")) >= 2 and float(row(pg, 1).locator("td").first.evaluate("e=>getComputedStyle(e).borderTopWidth").replace("px", "")) <= 1 and row(pg, 0).locator("td").first.evaluate("e=>getComputedStyle(e).borderTopColor") == "rgb(0, 0, 0)", f"{bgp(-4)} {bgp(1)}")
    check("العمود الأول غير ثابت في الطباعة (تأثير الهاتف مُلغى)", pg.evaluate("getComputedStyle(document.querySelector('#rp-table td')).position") == "static")
    f1 = f"{OUT}/rp_report.pdf"; pg.pdf(path=f1, prefer_css_page_size=True, print_background=False)
    rd = PdfReader(f1); t1 = ptext(f1, 1)
    check("PDF: A4 وصفحة واحدة لتسعة أيام", len(rd.pages) == 1 and abs(float(rd.pages[0].mediabox[2]) - 595.3) < 2, str(len(rd.pages)))
    check("PDF: كل تواريخ الفترة موجودة", all(fmt(iso(i)) in t1 for i in range(-4, 5)))
    check("PDF: لا عناصر تطبيق (أزرار/تنقل)", not any(x in t1 for x in ("عرض الكشف", "Duty Planner", "فتح اليوم", "me@test.com")))
    subprocess.run(["pdftoppm", "-r", "75", "-gray", "-png", f1, f"{OUT}/rp_report"], check=True)
    # الهاتف + الطباعة: لا تأثير للهاتف
    pg.set_viewport_size({"width": 390, "height": 800})
    check("طباعة بعرض هاتف: يبقى جدولًا ثابتًا بلا تمرير أو بطاقات", pg.evaluate("[getComputedStyle(document.querySelector('#rp-table')).display,getComputedStyle(document.querySelector('#rp-table')).minWidth,getComputedStyle(document.querySelector('#rp-table td')).position]") == ["table", "0px", "static"])
    pg.emulate_media(media="screen"); pg.set_viewport_size({"width": 1280, "height": 900})

    # ===== 9) فترة طويلة: كل الأيام في الطباعة (ليس الصفحة الأولى فقط) =====
    period(pg, iso(-30), iso(29)); apply(pg)
    check("فترة 60 يومًا: 60 صفًا", rows(pg).count() == 60 and "60 يومًا" in pg.inner_text("#rp-days"))
    pg.emulate_media(media="print"); pg.set_viewport_size({"width": 794, "height": 1123})
    f2 = f"{OUT}/rp_long.pdf"; pg.pdf(path=f2, prefer_css_page_size=True, print_background=False); n2 = len(PdfReader(f2).pages)
    alltxt = "\n".join(ptext(f2, i) for i in range(1, n2 + 1))
    check("PDF الطويل: صفحات متعددة", n2 >= 2, str(n2))
    check("PDF الطويل: كل التواريخ الستين مرة واحدة كحد أدنى (لا قصّ للصفحة الأولى)", all(fmt(iso(i)) in alltxt for i in range(-30, 30)) and fmt(iso(29)) in ptext(f2, n2))
    check("PDF الطويل: رأس الجدول يتكرر في كل صفحة", all(re.search(r"2026", ptext(f2, i)) for i in range(1, n2 + 1)) and ptext(f2, 2).count("التاريخ") + ptext(f2, 2).count("ﺦﻳﺭﺎﺘﻟﺍ") + ptext(f2, 2).count("الخفارة") + ptext(f2, 2).count("ﺓﺭﺎﻔﺨﻟﺍ") >= 1 or "ﺓﺭﺎﻔﺨﻟﺍ" in ptext(f2, 2) or "ةراﻔﺨﻟا" in ptext(f2, 2))
    subprocess.run(["pdftoppm", "-r", "60", "-gray", "-png", "-f", "1", "-l", "2", f2, f"{OUT}/rp_long"], check=True)
    pg.emulate_media(media="screen"); pg.set_viewport_size({"width": 1280, "height": 900})

    # ===== 10) تصدير Word حقيقي =====
    period(pg, iso(-4), iso(4)); apply(pg)
    with pg.expect_download() as dl: pg.click("[data-action=export-word]")
    d = dl.value; path = f"{OUT}/rp_export.docx"; d.save_as(path)
    check("اسم الملف .docx ويحمل الفترة", d.suggested_filename == f"كشف-الخدمات-{iso(-4)}_{iso(4)}.docx", d.suggested_filename)
    head = open(path, "rb").read(4)
    z = zipfile.ZipFile(path)
    check("الملف DOCX حقيقي (ZIP بأجزاء Word) وليس HTML أو TXT", head == b"PK\x03\x04" and z.testzip() is None and {"[Content_Types].xml", "word/document.xml", "word/styles.xml", "_rels/.rels"} <= set(z.namelist()) and b"<html" not in z.read("word/document.xml").lower())
    for n in z.namelist(): etree.fromstring(z.read(n))
    check("كل أجزاء XML سليمة", True)
    doc = pydocx.Document(path); tbl = doc.tables[1]
    check("python-docx يفتح الملف: جدول الكشف بعشرة صفوف (رأس + 9 أيام) وخمسة أعمدة", len(tbl.rows) == 10 and len(tbl.columns) == 5)
    check("رأس الجدول", [c.text for c in tbl.rows[0].cells] == ["التاريخ واليوم", "الخفارة", "الحراسة", "الراحة", "المهام"])
    check("كل أيام الفترة بالترتيب", [tbl.rows[i + 1].cells[0].text.split("\n")[0] for i in range(9)] == [fmt(iso(i)) for i in range(-4, 5)])
    check("الأسماء في خلاياها وكل اسم في فقرة مستقلة", [c.text for c in tbl.rows[1].cells][1:4] == ["بن ناصر", "عقون", "بوثلجة\nحميدي\nدبزة"] and "مرزوقي 🚗 — حراسة المخزن" in tbl.rows[1].cells[4].text and "لا توجد" in tbl.rows[4].cells[1].text)
    check("البديل بنوعه، وغير النشط معلَّم", "بديل: بن ناصر" in tbl.rows[2].cells[1].text and "شيباني (غير نشط)" in tbl.rows[9].cells[3].text)
    check("حالات الأيام في الملف (منفذ/قادم/اليوم/غير محسوم/غير مُعدّ)", "✓ منفذ" in tbl.rows[1].cells[0].text and "مسودة غير محسومة" in tbl.rows[3].cells[0].text and "غير مُعدّ" in tbl.rows[4].cells[0].text and "اليوم" in tbl.rows[5].cells[0].text and "قادم (معتمد)" in tbl.rows[6].cells[0].text and "منفذ" not in tbl.rows[6].cells[0].text)
    txt = "\n".join(p_.text for p_ in doc.paragraphs)
    check("العنوان والفترة وتاريخ الإعداد وسطر التعداد", "كشف الخدمات اليومية" in txt and f"الفترة: من {fmt(iso(-4))} إلى {fmt(iso(4))}" in txt and f"تاريخ إعداد التقرير: {fmt(T.isoformat())}" in txt and "9 أيام" in txt)
    check("وسيلة الإيضاح في الملف (5 خلايا)", len(doc.tables[0].rows[0].cells) == 5 and "منفذ" in doc.tables[0].rows[0].cells[0].text and "غير مُعدّ" in doc.tables[0].rows[0].cells[4].text)
    xml = z.read("word/document.xml").decode("utf8")
    check("اتجاه RTL: bidi على الفقرات والقسم وbidiVisual على الجدول وrtl على المقاطع", xml.count("<w:bidi/>") > 50 and "<w:bidiVisual/>" in xml and xml.count("<w:rtl/>") > 50 and re.search(r"<w:sectPr>.*<w:bidi/></w:sectPr>", xml))
    check("تظليل الحالات: منفذ (2 صفين×5 + الوسيلة)، غير محسوم (5+1)، اليوم بحدود سميكة", xml.count('w:fill="EEF0F2"') == 11 and xml.count('w:fill="FFF7E3"') == 6 and xml.count('w:sz="18"') >= 10)
    check("الصف لا ينقسم بين صفحتين ورأس الجدول يتكرر", xml.count("<w:cantSplit/>") >= 10 and xml.count("<w:tblHeader/>") == 1)
    lo = f"{OUT}/lo"; os.makedirs(lo, exist_ok=True)
    r_ = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", lo, path], capture_output=True, text=True, timeout=150)
    lp = f"{lo}/rp_export.pdf"
    check("LibreOffice يفتح الملف ويحوّله إلى PDF", os.path.exists(lp) and len(PdfReader(lp).pages) >= 1, r_.stderr[-120:])
    check("نص الملف بعد الفتح يحوي كل التواريخ التسعة", all(fmt(iso(i)) in ptext(lp, 1) for i in range(-4, 5)))
    subprocess.run(["pdftoppm", "-r", "75", "-png", lp, f"{OUT}/rp_export"], check=True)
    # تصدير مفلتر وطويل
    pg.fill("#rp-search", "بن ناصر")
    with pg.expect_download() as dl: pg.click("[data-action=export-word]")
    dl.value.save_as(f"{OUT}/rp_filtered.docx"); df = pydocx.Document(f"{OUT}/rp_filtered.docx")
    check("التصدير المفلتر: 5 أيام + سطر التصفية في الملف", len(df.tables[1].rows) == 6 and any("تصفية مطبّقة" in p_.text and "بن ناصر" in p_.text for p_ in df.paragraphs))
    pg.click("#rp-reset"); period(pg, iso(-30), iso(29)); apply(pg)
    with pg.expect_download() as dl: pg.click("[data-action=export-word]")
    dl.value.save_as(f"{OUT}/rp_long.docx"); dlg = pydocx.Document(f"{OUT}/rp_long.docx")
    check("تصدير 60 يومًا: كل الأيام في الملف (61 صفًا) لا ما يظهر على الشاشة فقط", len(dlg.tables[1].rows) == 61 and dlg.tables[1].rows[60].cells[0].text.startswith(fmt(iso(29))))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", lo, f"{OUT}/rp_long.docx"], capture_output=True, timeout=150)
    check("LibreOffice: الملف الطويل صفحات متعددة بكل التواريخ", len(PdfReader(f"{lo}/rp_long.pdf").pages) >= 2 and all(fmt(iso(i)) in "".join(ptext(f"{lo}/rp_long.pdf", k) for k in range(1, len(PdfReader(f"{lo}/rp_long.pdf").pages) + 1)) for i in range(-30, 30)))

    # ===== 11) سلامة عامة =====
    check("لا شيء كُتب في Firebase طوال الاختبار", raw(pg) == base_raw)
    check("لا مجموعات جديدة (لا statistics)", {k.split("/")[0] for k in store(pg)} <= {"people", "assignments", "daily_services", "settings"})
    check("الأفراد الـ18 والأيام السبعة والتعيينان لم تُمسّ", len([k for k in store(pg) if k.startswith("people/")]) == 18 and len([k for k in store(pg) if k.startswith("daily_services/")]) == 7 and len([k for k in store(pg) if k.startswith("assignments/")]) == 2)
    for _ in range(3): go(pg, "people", ".person"); go(pg, "stats", "#rp-table")
    pg.evaluate("()=>{window.__p={n:0};window.print=()=>{window.__p.n++}}"); pg.click("[data-action=print-report]")
    check("لا تراكم أحداث بعد زيارات متكررة (طباعة واحدة لنقرة واحدة)", pg.evaluate("window.__p.n") == 1)

    # ===== 12) الهاتف =====
    m = setup(b, 390, 844); login(m); go(m, "people", ".person")
    inject(m, {f"daily_services/{iso(0)}": D(iso(0), status="final", parmanoPersonId=pid(1), gardPersonId=pid(2), restPersonIds=[pid(3), pid(7), pid(8), pid(10)]), f"daily_services/{iso(1)}": D(iso(1), parmanoPersonId=pid(4))})
    go(m, "stats", "#rp-table"); m.fill("#rp-from", iso(-2)); m.fill("#rp-to", iso(5)); m.click("#rp-apply"); m.wait_for_selector("#rp-table tbody tr")
    check("الهاتف: الصفحة بلا تمرير أفقي والجدول يتمرر أفقيًا داخل حاويته", not m.evaluate("document.documentElement.scrollWidth>document.documentElement.clientWidth") and m.evaluate("(()=>{const w=document.querySelector('.rp-wrap');return w.scrollWidth>w.clientWidth&&getComputedStyle(w).overflowX==='auto'})()"))
    check("الهاتف: عمود التاريخ ثابت أثناء التمرير والأسماء كاملة", m.evaluate("getComputedStyle(document.querySelector('#rp-table tbody td')).position") == "sticky" and "شيتور" in m.inner_text("#rp-table"))
    check("الهاتف: أدوات الفترة والتصدير ظاهرة داخل الشاشة", all(m.evaluate("s=>{const r=document.querySelector(s).getBoundingClientRect();return r.left>=0&&r.right<=innerWidth&&r.width>0}", s) for s in ("#rp-apply", "[data-action=export-word]", "[data-action=print-report]", "#rp-search")))
    m.screenshot(path=f"{OUT}/rp_mobile.png")
    b.close()

print("\n--- أخطاء المتصفح ---"); print("\n".join(errors) if errors else "لا توجد أخطاء JavaScript")
print(f"\n{sum(results)}/{len(results)} نجحت"); sys.exit(0 if all(results) and not errors else 1)
